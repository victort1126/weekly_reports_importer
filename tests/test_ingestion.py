from __future__ import annotations

import unittest
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from file_ingest import (
    FileIngestionError,
    ingest_slack_file,
    ingest_slack_files,
    is_supported_report_file,
)
from markdown_format import convert_markdown
from models import EnrichedReport, ReportBatch, WeeklyReport


class _Response:
    def __init__(self, content: bytes) -> None:
        self.content = content
        self.headers = {"Content-Length": str(len(content))}

    def raise_for_status(self) -> None:
        return None

    def iter_content(self, chunk_size: int):
        del chunk_size
        yield self.content


class IngestionTests(unittest.TestCase):
    @staticmethod
    def _workbook_bytes(name: str, task: str) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name", "Completed"])
        sheet.append([name, task])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()

    def test_emojis_are_removed_but_markdown_and_lines_remain(self) -> None:
        report = WeeklyReport(
            employee_name="Jane 👩‍💻",
            team="Platform 🚀",
            iso_year_week="2026-W30",
            tasks_completed="📌 *Completed*\n- shipped feature ✅",
        )

        self.assertEqual(report.employee_name, "Jane")
        self.assertEqual(report.team, "Platform")
        self.assertEqual(report.tasks_completed, "*Completed*\n- shipped feature")

        plain, spans = convert_markdown(report.tasks_completed)
        self.assertEqual(plain, "Completed\n- shipped feature")
        self.assertTrue(spans[0].bold)

    def test_enriched_roster_values_are_also_cleaned(self) -> None:
        report = EnrichedReport(
            iso_year_week="2026-W30",
            employee_name="Jane 🚀",
            team="Platform ✨",
            tasks_completed="Done ✅",
        )
        self.assertEqual(report.to_row()[1:4], ["Jane", "Platform", "Done"])

    def test_markdown_file_download_and_decode(self) -> None:
        content = "# Weekly\n\n*Done*\n- task A".encode()
        file = {
            "name": "weekly.md",
            "url_private_download": "https://files.slack.test/weekly.md",
        }

        with patch("file_ingest.requests.get", return_value=_Response(content)):
            result = ingest_slack_file(file, "xoxb-test")

        self.assertEqual(result.name, "weekly.md")
        self.assertEqual(result.text, "# Weekly\n\n*Done*\n- task A")

    def test_xlsx_bundle_is_rendered_with_all_worksheets_and_rows(self) -> None:
        workbook = Workbook()
        first = workbook.active
        first.title = "Platform"
        first.append(["Name", "Completed", "Next"])
        first.append(["Jane Doe", "Feature A", "Feature B"])
        second = workbook.create_sheet("Growth")
        second.append(["Name", "Completed", "Next"])
        second.append(["John Smith", "Campaign A", "Campaign B"])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        file = {
            "name": "team-weekly.xlsx",
            "url_private_download": "https://files.slack.test/team-weekly.xlsx",
        }
        with patch(
            "file_ingest.requests.get", return_value=_Response(buffer.getvalue())
        ):
            result = ingest_slack_file(file, "xoxb-test")

        self.assertIn("=== Worksheet: Platform ===", result.text)
        self.assertIn("Name (employee_name): Jane Doe", result.text)
        self.assertIn("Completed (tasks_completed): Feature A", result.text)
        self.assertIn("=== Worksheet: Growth ===", result.text)
        self.assertFalse(result.needs_llm)
        self.assertEqual(
            [report.employee_name for report in result.reports],
            ["Jane Doe", "John Smith"],
        )
        self.assertEqual(result.reports[0].tasks_completed, "Feature A")
        self.assertEqual(result.reports[0].next_steps, "Feature B")
        self.assertEqual(result.reports[1].tasks_completed, "Campaign A")

    def test_common_xslx_filename_typo_is_accepted(self) -> None:
        self.assertTrue(is_supported_report_file({"name": "bundle.xslx"}))

    def test_zip_bundle_produces_one_source_per_workbook(self) -> None:
        archive_buffer = BytesIO()
        with ZipFile(archive_buffer, "w", ZIP_DEFLATED) as archive:
            archive.writestr(
                "Platform/Jane Doe.xlsx",
                self._workbook_bytes("Jane Doe", "Feature A"),
            )
            archive.writestr(
                "Growth/John Smith.xslx",
                self._workbook_bytes("John Smith", "Campaign A"),
            )
            archive.writestr("__MACOSX/._junk", b"ignored")
            archive.writestr("notes.txt", b"ignored")

        file = {
            "name": "team-weeklies.zip",
            "url_private_download": "https://files.slack.test/team-weeklies.zip",
        }
        with patch(
            "file_ingest.requests.get",
            return_value=_Response(archive_buffer.getvalue()),
        ):
            results = ingest_slack_files(file, "xoxb-test")

        self.assertEqual(len(results), 2)
        self.assertEqual(
            [result.name for result in results],
            [
                "team-weeklies.zip / Platform/Jane Doe.xlsx",
                "team-weeklies.zip / Growth/John Smith.xslx",
            ],
        )
        self.assertIn("Jane Doe", results[0].text)
        self.assertIn("Feature A", results[0].text)
        self.assertEqual(results[0].reports[0].employee_name, "Jane Doe")
        self.assertFalse(results[0].needs_llm)
        self.assertIn("John Smith", results[1].text)
        self.assertIn("Campaign A", results[1].text)
        self.assertEqual(results[1].reports[0].employee_name, "John Smith")
        self.assertFalse(results[1].needs_llm)

    def test_zip_with_no_supported_reports_is_rejected(self) -> None:
        archive_buffer = BytesIO()
        with ZipFile(archive_buffer, "w", ZIP_DEFLATED) as archive:
            archive.writestr("readme.txt", b"nothing to import")

        file = {
            "name": "empty.zip",
            "url_private_download": "https://files.slack.test/empty.zip",
        }
        with (
            patch(
                "file_ingest.requests.get",
                return_value=_Response(archive_buffer.getvalue()),
            ),
            self.assertRaisesRegex(
                FileIngestionError, "contains no supported"
            ),
        ):
            ingest_slack_files(file, "xoxb-test")

    def test_zip_is_recognized_as_supported(self) -> None:
        self.assertTrue(is_supported_report_file({"name": "team-weeklies.zip"}))

    def test_report_batch_drops_blank_tasks_instead_of_failing(self) -> None:
        """A stub first report (the original Excel failure) must not abort the batch."""
        batch = ReportBatch.model_validate(
            {
                "reports": [
                    {
                        "employee_name": "Header Row",
                        "iso_year_week": "2026-W32",
                        "tasks_completed": "",
                    },
                    {
                        "employee_name": "王怡明",
                        "iso_year_week": "2026-W32",
                        "tasks_completed": "完成跌倒识别",
                    },
                ]
            }
        )
        self.assertEqual([report.employee_name for report in batch.reports], ["王怡明"])

    def test_report_batch_keeps_employees_separate(self) -> None:
        batch = ReportBatch(
            reports=[
                WeeklyReport(
                    employee_name="Jane",
                    iso_year_week="2026-W30",
                    tasks_completed="A",
                ),
                WeeklyReport(
                    employee_name="John",
                    iso_year_week="2026-W30",
                    tasks_completed="B",
                ),
            ]
        )
        self.assertEqual([report.employee_name for report in batch.reports], ["Jane", "John"])


if __name__ == "__main__":
    unittest.main()
